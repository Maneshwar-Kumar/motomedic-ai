"""Insert the requested guardrail sections into the two existing planning DOCX files."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent


def insert_after(anchor, text: str, style: str):
    document = anchor._parent
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)
    anchor._p.addnext(paragraph._p)
    return paragraph


def find_paragraph(document: Document, exact: str):
    return next((p for p in document.paragraphs if p.text.strip() == exact), None)


def replace_text(paragraph, old: str, new: str) -> None:
    if old not in paragraph.text:
        return
    if len(paragraph.runs) == 1:
        paragraph.runs[0].text = paragraph.runs[0].text.replace(old, new)
    else:
        # These targeted paragraphs have a single text run in the source files.
        raise ValueError(f"Unexpected split runs in: {paragraph.text}")


def refresh_manual_status(doc: Document, is_tracker: bool = False) -> None:
    for paragraph in doc.paragraphs:
        if "Manual-specific rules and golden answers pending OEM source review" in paragraph.text:
            replace_text(paragraph, "Manual-specific rules and golden answers pending OEM source review", "Owner-manual prototype; independent review and variant fit pending")
        if "An OEM manual with enough troubleshooting structure and variant-specific procedures can be sourced and lawfully used" in paragraph.text:
            replace_text(paragraph, paragraph.text, "The supplied owner's manual is available for this prototype; confirm variant fit and rights before a public demo")
        if is_tracker:
            replacements = {
                "4  Golden-answer dataset plan: deferred until manual arrives": "4  Golden-answer dataset: provisional atomic cases",
                "Do not invent OEM facts. After the manual is uploaded, first assess manual suitability and annotate authoritative spans. Then draft 25–100 labelled examples with multiple turns and photograph scenarios only where the manual supports them.": "Do not invent OEM facts. The supplied owner's manual has been indexed provisionally, and 37 atomic dialogue and lookup cases are drafted with source pages. Independently review their spans, outcomes and variant fit before approval; expand toward 50–100 cases, adding photograph scenarios only where evidence supports them.",
                "34 atomic dialogue and lookup cases": "37 atomic dialogue and lookup cases",
                "Without the actual manual, neither a credible golden-answer corpus nor model-specific performance estimates exist": "The owner manual supports a provisional source-checked corpus; independent review, variant confirmation and held-out evaluation are still required before performance claims",
            }
            for old, new in replacements.items():
                if old in paragraph.text:
                    replace_text(paragraph, old, new)
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            replace_text(paragraph, "PRE-MANUAL", "OWNER-MANUAL PILOT")
    if is_tracker:
        for table in doc.tables:
            for row in table.rows:
                label = row.cells[0].text
                if label.startswith("Golden Q&A set"):
                    row.cells[1].text = "DRAFT: 37 atomic cases; human review pending"
                elif label.startswith("Source, license and configuration manifest"):
                    row.cells[1].text = "DRAFT: verify variant and rights"


def concept() -> None:
    path = ROOT / "01_Concept_Architecture.docx"
    doc = Document(path)
    refresh_manual_status(doc)
    if find_paragraph(doc, "Guardrails and release checks"):
        doc.save(path)
        print(f"Refreshed: {path.name}")
        return
    anchor = next(p for p in doc.paragraphs if "If retrieval infrastructure is unavailable" in p.text)
    anchor = insert_after(anchor, "Guardrails and release checks", "Heading 1")
    items = [
        "Source and applicability gate: confirm the selected owner manual checksum, model and variant before any procedure; block advice when applicability is unresolved.",
        "Evidence gate: assemble the relevant procedure, conditions, warnings, specifications, tables and linked pages; require every technical claim to map to an immutable source object and page.",
        "Safety gate: surface source warnings and unmet prerequisites, stop unsafe or undocumented bypass requests, and refer to authorised service when the manual gives no safe next step.",
        "Dialogue guardrail: ask one discriminating, source-derived question at a time; store typed outcomes; never infer that an attempted action was normal or repeat it without a documented reason.",
        "Untrusted input guardrail: treat PDF text, images and user messages as data; ignore embedded instructions that try to change source, safety or citation policy.",
        "Failure distinction: insufficient evidence triggers clarification or honest abstention; a parser, retrieval or model outage triggers a system-error message, never a false claim that the manual lacks information.",
        "Cache guardrail: cache exact reviewed golden paths by manual checksum, variant, flow and policy version. Invalidate on any relevant change, and never use similarity alone to select a cached repair answer.",
        "Privacy and cost guardrail: keep photos and audio ephemeral by default, redact traces, protect secrets, and apply rate and spending limits to a public demo.",
    ]
    for item in items:
        anchor = insert_after(anchor, item, "List Bullet")
    doc.save(path)
    print(f"Updated: {path.name}")


def to_build() -> None:
    path = ROOT / "04_To_Build_Evaluation_Decisions.docx"
    doc = Document(path)
    refresh_manual_status(doc, is_tracker=True)
    if find_paragraph(doc, "11  Guardrail implementation checklist"):
        doc.save(path)
        print(f"Refreshed: {path.name}")
        return
    anchor = next(p for p in doc.paragraphs if "A polished hosted demo is necessary" in p.text)
    anchor = insert_after(anchor, "11  Guardrail implementation checklist", "Heading 1")
    items = [
        "G-01  Lock source checksum, manual edition and confirmed bike variant in each session; reject mismatches before retrieval or cached-answer lookup.",
        "G-02  Check extraction completeness for warnings, prerequisites, table headers/units and figure references; send uncertain pages to manual review.",
        "G-03  Gate every proposed technical step on applicability, complete evidence, relevant cautions and claim-level citation support; reject unsupported drafts.",
        "G-04  Implement deterministic typed diagnostic transitions, one question at a time, with no unjustified repeat after 'tried but still failing'.",
        "G-05  Separate supported answer, clarification, unsupported request, safety stop and system error in both UI and telemetry.",
        "G-06  Test prompt injection in user text, uploaded images and the manual; none may override source or safety rules.",
        "G-07  Compile approved golden dialogue prefixes and terminal answers into exact-match cache keys containing manual checksum, variant, flow and policy version; fail closed on stale or draft entries.",
        "G-08  Keep raw media out of ordinary logs, store keys server-side, and set file, rate and spend limits before a public demo.",
        "G-09  Add regression cases for wrong variant, missing warning, ambiguous visual, invalid citation, retrieval outage, exhausted paths and cache mismatch; require a review sign-off before release.",
    ]
    for item in items:
        anchor = insert_after(anchor, item, "List Bullet")
    table = doc.tables[0]
    row = table.add_row()
    for cell, value in zip(row.cells, [
        "Guardrail implementation and safety regression suite",
        "TO BUILD; design added in section 11",
        "All G-01 through G-09 checks pass against reviewed cases, including cache invalidation and error-state tests",
    ]):
        cell.text = value
    doc.save(path)
    print(f"Updated: {path.name}")


if __name__ == "__main__":
    concept()
    to_build()
