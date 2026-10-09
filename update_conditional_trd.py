"""Apply the agreed conditional retrieval policy to the existing TRD."""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt


def insert_after(anchor, text, style="Normal"):
    paragraph = anchor._parent.add_paragraph(text, style)
    anchor._p.addnext(paragraph._p)
    return paragraph


def polish(doc):
    # Cell.text resets direct run formatting; restore the existing 8-point table style.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8)
        if table.rows[-1].cells[0].text == "RetrievalDecision":
            for target, donor in zip(table.rows[-1].cells, table.rows[-3].cells):
                current = target._tc.tcPr
                if current is not None:
                    target._tc.remove(current)
                target._tc.insert(0,deepcopy(donor._tc.tcPr))
    for section in doc.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in footer.paragraphs:
                for run in paragraph.runs:
                    if "v0.3" in run.text:
                        run.text = run.text.replace("v0.3","v0.4")


def update(path):
    doc = Document(path)
    if any(p.text == "4 2 Conditional reranking and bounded recovery" for p in doc.paragraphs):
        polish(doc)
        doc.save(path)
        return
    replacements = {
        "TRD v0.3  Reference design and implementation contracts": "TRD v0.4  Reference design and implementation contracts",
        "DRAFT v0.3  Updated 9 October 2026  Independent source review pending": "DRAFT v0.4  Updated 9 October 2026  Independent source review pending",
        "2  Offline ingestion pipeline": "2  Versioned hosted ingestion pipeline",
    }
    for paragraph in doc.paragraphs:
        if paragraph.text in replacements:
            paragraph.text = replacements[paragraph.text]
        elif paragraph.text.startswith("First-stage dense retrieval uses"):
            paragraph.text = "First-stage dense retrieval uses cosine similarity or the embedding model's recommended metric. Sparse BM25 retrieval handles exact warning labels, component codes and units. Fuse ranks with Reciprocal Rank Fusion, assemble initial parent and verified dependency context, and inspect ranking uncertainty separately from source applicability and evidence coverage. Conditionally rerank a bounded candidate pool before reassembly and the controller/answer gates."
        elif paragraph.text.startswith("Candidate settings are starting hypotheses"):
            paragraph.text = "The hosted prototype uses a fused pool of 20 candidates and selects five hits for assembly. One recovery may widen the fused pool to 40 and fetch exact reviewed dependency IDs. Reranking scores at most 20 candidates and 100 token-bounded query/passage pairs. At most 256 mandatory dependency IDs may be attached; an over-limit manifest fails as SYSTEM_ERROR and must not be silently truncated. These limits and trigger thresholds are versioned starting hypotheses, to be calibrated on independently reviewed cases."
        elif paragraph.text.startswith("If evidence remains incomplete, permit"):
            paragraph.text = "If an approved dependency set is incomplete, permit one bounded recovery that fetches the exact required objects and widens retrieval once. If dependencies remain unavailable, withhold technical advice and record a non-rider-resolvable source gap. Allow at most one subsequent reranking attempt when ordering remains uncertain and all source/applicability checks permit it; never loop between search, reranking and generation."
        elif paragraph.text.startswith("Initial benchmarks: text warm"):
            paragraph.text += " Measure time to first output (TTFO) from confirmed rider input to the first releasable answer token. A loading indicator is excluded. Technical content is withheld until the required claim and citation checks pass. Report fast-path and recovery/reranking p50/p95, their combined p95, fallback frequency and cold starts separately; reduced average latency alone does not establish a lower p95."
        elif paragraph.text.startswith("Version together: manual checksum"):
            paragraph.text = paragraph.text.replace("controller policy, reranker,", "controller policy, retrieval policy and trigger thresholds, calibration dataset and approval, reranker,")
    anchor = next(p for p in doc.paragraphs if p.text.startswith("Perform the existing bounded recovery search"))
    paragraphs = [
        ("4 2 Conditional reranking and bounded recovery", "Heading 2"),
        ("The normal request path uses conditional reranking. Clearly ordered candidates with complete reviewed dependencies continue to the eligibility and answer gates. Reranking handles uncertain ordering among plausible candidates; evidence recovery handles required objects absent from the initial bundle. Similarity, RRF and cross-encoder scores cannot prove support or approve a procedure.", "Normal"),
        ("Trigger signals compare distinct parent sections: disagreement between dense and BM25 top sections, or a narrow relative RRF margin between the two leading sections. Sibling chunks from the same section do not create ambiguity. Zero BM25 matches provide no sparse-agreement signal. Initial experimental values are top-three section overlap below 0.5 or relative RRF margin at most 0.05. Mark these values PROVISIONAL_UNCALIBRATED; use reviewed tuning cases to calibrate them and held-out cases to assess their effect before a release decision.", "Normal"),
        ("Apply the routing policy in order. A known safety stop retains its verified warning and skips model work; a known system failure remains SYSTEM_ERROR; a known unsupported request abstains. Missing, stale or unapproved source/dependency review enters source review or abstention. Unconfirmed applicability enters one focused clarification; a confirmed mismatch blocks the action. An approved missing dependency triggers one exact-ID fetch and bounded retrieval expansion. Complete eligible evidence with uncertain ordering triggers one cross-encoder attempt; otherwise continue through the remaining gates.", "Normal"),
        ("Reassemble after reranking and recheck the dependency set. Retain all already fetched mandatory action, warning, prerequisite, table/header/unit and figure objects for the selected reviewed step even when candidate ordering changes. Reranking does not select a new diagnostic step or infer an OEM order. The controller remains responsible for selecting an eligible step and the final verifier remains responsible for claim support.", "Normal"),
        ("Operational failures in embedding, retrieval, assembly or the cross-encoder yield SYSTEM_ERROR with a typed error class and recorded attempts. A query exceeding the current cross-encoder's 128-token input bound, or work exceeding 100 query/passage pairs, skips that inference with an explicit reason and keeps ranking unresolved. Preserve the full query and source text; never silently truncate a rider report or mandatory evidence. A source gap cannot be resolved by asking the rider to supply a missing manual warning or table header.", "Normal"),
        ("Return a RetrievalDecision with policy version/configuration, calibration status, execution mode, uncertainty signals, decision history, recovery/rerank attempt counts, exact missing/fetched evidence IDs, pair count, skip reason and stage durations. Conditional mode is the normal code path; off and always_diagnostic modes exist only for ablations. Always-diagnostic execution can bypass the scheduling decision but cannot bypass evidence, applicability, safety or answer-release gates. A generic insufficient-bundle flag must never trigger reranking by itself.", "Normal"),
        ("The current hosted manual has pending applicability and independent dependency review. Its candidate bundles therefore skip conditional reranking with explicit review/clarification reasons and keep may_generate_procedural_advice false. The implementation is available for approved registry/controller inputs; source and release approvals are separate work. Changing this policy versions retrieval and invalidates affected response caches while reusing unchanged document embeddings.", "Normal"),
        ("Evaluate off, always and conditional policies on the same independently reviewed cases. Record complete evidence recall, rescued and regressed cases, proper abstention, wrong-variant leakage, trigger precision/recall, reranking frequency, per-stage latency, TTFO and cost per correctly handled case. Adopt calibrated conditional reranking for release only if it improves reviewed outcomes without violating latency/cost and safety gates. The existing 27 draft in-sample page-coverage probes cannot establish that result.", "Normal"),
        ("Cross-encoder retrieval reference: https://www.sbert.net/examples/sentence_transformer/applications/retrieve_rerank/README.html", "Normal"),
    ]
    for text, style in paragraphs:
        anchor = insert_after(anchor,text,style)
    for table in doc.tables:
        header = table.rows[0].cells[0].text
        for row in table.rows[1:]:
            if row.cells[0].text == "Conditional reranker":
                row.cells[1].text = "One attempt for uncertain candidate ordering after source and coverage checks"
                row.cells[2].text = "Versioned provisional thresholds; reviewed calibration; bounded pool; mandatory evidence retained"
            elif row.cells[0].text == "rerank / graph_expand":
                row.cells[1].text = "policy_version, calibration_status, mode, trigger_reasons, ranked_ids, pair_count, skip_reason, verified_edges, added_dependencies, attempt_counts, duration_ms"
            elif row.cells[0].text == "E3":
                row.cells[1].text = "Off vs always vs conditional cross-encoder; optional late-interaction alternative"
                row.cells[2].text = "Complete evidence recall, rescues/regressions, trigger accuracy, fallback rate, TTFO p95, cost"
        if header == "Entity":
            cells = table.add_row().cells
            cells[0].text = "RetrievalDecision"
            cells[1].text = "policy_version/config, calibration_status, mode, uncertainty_signals, decisions, recovery_attempts, rerank_attempts, requested/fetched/missing_evidence_ids, pair_count, skip_reason, stages_seconds, final_action"
    polish(doc)
    doc.save(path)


if __name__ == "__main__":
    update(Path(__file__).with_name("03_TRD.docx"))
