"""Structural document QA, explicitly distinct from rendered all-page inspection."""
from pathlib import Path
import hashlib,json,zipfile
from docx import Document
from lxml import etree
from tools.reconcile_draft_verify_stream_documents import NAMES,MAPS,OUT

def main():
 records=[]
 for name in NAMES:
  path=Path(name);prior=OUT/'prior_documents'/name;doc=Document(path);old=Document(prior);skip=set(MAPS.get(name,{}))
  if name=='03_TRD_Architecture_Decision_Note.docx':skip.update({1,2});skip.update(i for i,p in enumerate(old.paragraphs) if p.style.name.startswith('Heading') and i>=3)
  if name=='04_To_Build_Evaluation_Decisions.docx':skip.update(i for i,p in enumerate(old.paragraphs) if p.text in ('Reconciliation change record 9 October 2026','Rider UI and safeguards reconciliation 11 October 2026'))
  retained=[p.text for i,p in enumerate(old.paragraphs) if i not in skip and p.text];missing=[p for p in retained if p not in [x.text for x in doc.paragraphs]]
  texts='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
  required=['DRAFT','VERIFY','STREAM','separate','actual question','partial','operation ID','durable save','ACTION_COMPLETED_UNRESOLVED','SYSTEM_ERROR','No predefined question','private drafting tokens','at most one','not authorize new spending']
  absent=[x for x in required if x not in texts]
  with zipfile.ZipFile(path) as z:
   for n in z.namelist():
    if n.endswith('.xml'):etree.fromstring(z.read(n))
   media={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
  with zipfile.ZipFile(prior) as z:oldmedia={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
  result={'file':name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'prior_sha256':hashlib.sha256(prior.read_bytes()).hexdigest(),'unrelated_existing_nonempty_paragraphs_checked':len(retained),'missing_unrelated_paragraphs':missing,'missing_current_requirements':absent,'original_tables_preserved':len(doc.tables)==len(old.tables),'original_media_byte_identical':media==oldmedia,'ooxml_parse_ok':True,'structural_result':'PASS' if not missing and not absent and media==oldmedia else 'FAIL','visual_qa':'PENDING_MANAGED_RENDER_AND_ALL_PAGE_INSPECTION'}
  if name=='03_TRD_Architecture_Decision_Note.docx':result['historical_numeric_comparison_table_content_identical']=[[c.text for c in r.cells] for r in doc.tables[1].rows]==[[c.text for c in r.cells] for r in old.tables[1].rows]
  if name=='02_PRD.docx':
   ids=[r.cells[0].text for r in doc.tables[1].rows[1:]];result['unique_functional_ids']=len(ids)==len(set(ids))
  records.append(result)
 passed=all(x['structural_result']=='PASS' for x in records)
 (OUT/'structural_qa.json').write_text(json.dumps({'checks':records,'all_structural_checks_pass':passed,'not_visual_acceptance':True},indent=2),encoding='utf-8');print(json.dumps(records,indent=2));return 0 if passed else 1
if __name__=='__main__':raise SystemExit(main())
