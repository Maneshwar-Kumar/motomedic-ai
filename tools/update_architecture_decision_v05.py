"""Record the owner's 11 October 2026 A/B decision without changing runtime."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
REPORT = json.loads((ROOT/'reports/architecture_owner_measurement_summary_v4.json').read_text())
NOTE = '03_TRD_Architecture_Decision_Note.docx'

def paragraph(doc, text, style=None):
    return doc.add_paragraph(text, style)

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,h in zip(t.rows[0].cells, headers): c.text=h
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); trpr.append(no)
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[ci])
            props=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            props.append(borders)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F4F6F8' if ri%2==0 else 'FFFFFF'));props.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(3)
                for r in p.runs:r.font.size=Pt(9);r.bold=(ri==0);r.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph()
    return t

def heading(doc,text): doc.add_heading(text,1)

def main():
    d=Document(ROOT/'03_TRD.docx')
    replacements={
        1:'TRD v0.5  Reference design and implementation contracts',
        2:'v0.5  Updated 11 October 2026  Owner decision selects A and B with streaming for A',
        3:'TRD v0.5 is the authoritative design and technical contract from 11 October 2026. It preserves v0.4 ingestion, retrieval, eligibility and safety contracts and adopts the owner-selected A/B answer methods in section 14. The approved owner manual governs motorcycle facts. Exact source approval scopes, applicability, usage rights and release gates remain binding; a methodology decision is not a public deployment approval.',
        35:'Reassemble after reranking and recheck the dependency set. Retain mandatory action, warning, prerequisite, table/header/unit and figure objects for the selected reviewed step. Reranking does not select a diagnostic step or infer OEM order. The controller selects an eligible step; final claim validation follows the A or B method in sections 6 and 14.',
        38:'Candidate bundles without approved applicability and dependency scopes remain blocked. Owner-delegated coordinator approval now exists for 161 exact source-content/initial-response scopes and the selected Himalayan BS6 manual scope in release_manifest_v3.json. Consume its exact bindings; do not extend approval to changed objects, full adaptive transitions, general vision or deployment. Conditional retrieval scheduling remains versioned and reuses unchanged document embeddings.',
        55:'Apply the deterministic rules gate to every method: valid manual and variant, complete approved dependencies, satisfied prerequisites, required warnings, no explicit conflict and valid immutable citations. In A, the model selects indices from exact reviewed wording and the server validates the selection and releases only the bound wording; no free technical paraphrase or separate online semantic verifier is permitted in A. In B, the model composes a concise draft and a separate semantic support evaluator checks each proposed claim against its source spans before release. Reject output that fails the applicable method. A receives a post facto semantic verification audit before UI integration, as specified in section 14.',
        58:'SUPPORTED requires an eligible action or applicable factual answer, complete approved dependencies, valid citations, satisfied safety and prerequisite checks and successful method-specific claim validation. A requires exact reviewed wording and valid allowlisted selections; B requires deterministic checks plus successful semantic verification. Both release only after validation and durable session save. Rejected drafts remain withheld. Bounded operational generation/verification failures return SYSTEM_ERROR; established source gaps yield UNSUPPORTED. A post facto audit does not retroactively authorize unsafe output.',
        101:'Recorded implementation: hosted BGE M3 candidate index index_a3282c4d98cf54595e4ba168 has 532 children and 277 section parents; authenticated Cloud Run Qdrant restores immutable private vectors on cold start. The historical backend smoke passed 11/11. Isolated A/B/C experiments and verified delivery hooks are measured; the A/B production service and browser/UI delivery boundary require integration and separate validation.',
        102:'The exact 161 source-content/initial-response scopes have coordinator approval under owner delegation in release_manifest_v3.json, checksum b8f8c7b5a9e1a299f0a90eecbf10127779b62b35fc105f0cd587de433e235298. The 17-test approval consumer validates 187 bound files. Remaining limitations include full adaptive transitions, general vision, changed runtime objects and source-specific gaps outside those scopes. End-user PDF, excerpt and figure delivery remain prohibited by source_manifest.json; approved paraphrases and citation metadata are the UI content boundary.',
        104:'The deployed retrieval record retains the TRD v0.3 candidate assembly boundary. Repository v0.4 retrieval scheduling and isolated answer experiments do not establish a deployed rider service. This v0.5 document selects A/B and A streaming prospectively, with follow-up audit and hosting work in section 14. Progress references remain cloud/deployment_status.json, cloud/live_smoke.json, the exact approval package and the architecture v4 reports.'
    }
    for i,text in replacements.items(): d.paragraphs[i].text=text
    d.tables[0].rows[6].cells[2].text='A exact reviewed selection checks; B semantic claim verification; fail-closed delivery'
    d.tables[3].rows[1].cells[1].text='Release only after applicability, complete dependencies, warnings, citations and method-specific claim validation under A or B pass.'
    d.tables[4].rows[4].cells[2].text='Approved paraphrase and immutable citation metadata only; no end-user PDF, excerpt or crop under current display policy'
    d.add_page_break();heading(d,'14 Owner selected answer methods and delivery')
    paragraph(d,'Decision recorded 11 October 2026. Retain A and B as the application methodologies. Use verified streaming delivery for A. Exclude C from the current UI path because it failed the clock case in both passes and took excessive paced time in the successful cases. This is the owner decision based on a bounded engineering comparison, not a general model quality or p95 finding.')
    table(d,['Method','Selected behavior','Release requirement'],[
        ['A component selection','Select among exact reviewed wording; deliver by streaming','Validate exact indices, scope, dependencies and session; emit the complete approved unit after durable save'],
        ['B short generation','Retain concise source-grounded composition with a separate verifier','Withhold the whole draft until semantic verification, deterministic gates and durable save pass'],
        ['C verified independent blocks','Retain historical experiment only; do not wire into the current UI','No release route selected; no silent fallback from A/B to C']
    ],[1.15,2.35,3.0])
    paragraph(d,'Streaming for A means HTTP/SSE delivery of a complete verified answer unit. Do not expose raw selection JSON, model token fragments or partial cautions as repair guidance. Preserve terminal complete/error events and session ownership. The UI distinguishes provider TTFT from first safe output; a progress indicator is not TTFO. B may use transport events after its whole-answer gate, but the owner has specifically selected A streaming.')
    heading(d,'14 1 Evidence and latency contract')
    paragraph(d,'Hosted run 38122505406 at commit e910c86d2fb8a0a2f955bbea4fa4cbedbe64d82d used gemma-4-26b-a4b-it with minimal thinking on GEN003 mirrors, EXP003 signal cancellation and GEN026 clock settings. Nine initial turns and nine immediate repeats produced 8/9 SUPPORTED outcomes in each pass. All six A/B initial cases passed. C clock delivered no output in either pass. The run made 27 attempts with zero retries: 26 known HTTP200 results and one unknown status.')
    paragraph(d,'Report retrieval wall time, provider TTFT, request-relative TTFT, actual first verified output and full completion separately. Keep fresh-worker startup, benchmark cooldown and idle provider-admission waits distinct. Removing idle waits arithmetically is not an observed unpaced run. Actual receipt was measured on hosted localhost HTTP/SSE; public network and browser rendering remain unmeasured. Three cases cannot establish p95 or architecture-wide superiority.')
    paragraph(d,'The candidate cache remains bounded and ephemeral. Similarity only locates a candidate; exact source, variant, selected step, typed state/history and runtime approval bindings authorize reuse. Only reviewed wording is stored, fresh retrieval and gates remain mandatory and B retains verification on hits. Three A repeats hit; no B/C repeats hit. Concurrent dense/BM25 retrieval remains implemented, but three paired controls demonstrated no latency gain.')
    heading(d,'14 2 Work before UI integration')
    paragraph(d,'First, run the removed semantic verifier post facto on the exact A outputs from all three cases, with repeat provenance retained. Bind the audited wording, claims, warnings and citations to the approved source and existing run. Record valid acceptance, valid rejection, invalid schema and provider/system failure separately; do not regenerate the answers or infer approval from a schema failure. This task is a new audit, not permission to alter old results or widen A beyond approved wording.')
    paragraph(d,'Second, investigate current Cloud Run minimum-instance pricing for the actual region, CPU, memory and billing mode, including idle cost, free tier and existing budget impact. Check current startup CPU boost settings and pricing and propose a controlled benchmark. Fresh-worker readiness was 56.64 seconds excluding controls, including 38.25 seconds for dense model/index load; this is not measured Cloud Run request cold latency. Do not change deployment, minimum instances, boost or billing as part of research.')
    paragraph(d,'The companion 03_TRD_Architecture_Decision_Note.docx records selected and rejected answer methods, measured retrieval/cache/streaming controls and broader design alternatives, distinguishing completed comparisons from options merely considered. Preserve the v4 raw results and independent audit as evidence. The concept, PRD and evaluation documents dated 9 October retain historical v0.4 language; this v0.5 amendment takes precedence for A/B method-specific validation and A streaming.')
    d.save(ROOT/'03_TRD.docx')

    n=Document();sec=n.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=n.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    n.styles['Normal'].font.size=Pt(10.5);n.styles['Normal'].paragraph_format.space_after=Pt(7)
    n.styles['Title'].font.size=Pt(22);n.styles['Heading 1'].font.size=Pt(15)
    paragraph(n,'MotoMedic answer architecture decisions','Title')
    paragraph(n,'Companion note to TRD v0.5  11 October 2026','Subtitle')
    paragraph(n,'We will retain A and B for the application and deliver A through verified streaming. C will remain an experiment outside the current UI path. This note explains the owner decision, the design choices considered and the evidence used, so a reader can distinguish tested decisions from options that still need evaluation.')
    heading(n,'Selected answer methods')
    table(n,['Choice','What it does','Decision and reason'],[
        ['A component selection','Gemma selects indices into exact reviewed wording; server supplies the approved text','Selected with streaming. All three cases passed in both passes; exact candidate reuse avoided generation on repeats.'],
        ['B short generation','Gemma writes concise paraphrases, followed by a separate claim verifier','Retained as a methodology. All three cases passed in both passes; supports wording flexibility while retaining verification.'],
        ['C verified blocks','Generation streams independent blocks, each released after verification','Not selected for the UI. Clock failed in both passes; successful paced response times were high and earlier safe output was not demonstrated.']
    ],[1.15,2.65,3.3])
    paragraph(n,'A removes the separate online semantic-verifier call only because the output is selected from exact approved wording. It retains source applicability, approved dependency completeness, warnings, prerequisites, typed session eligibility, citation checks and durable session save. A is not an unrestricted generator without verification. B retains its semantic verifier. No model score can authorize a new repair instruction.')
    paragraph(n,'A streaming will emit the whole verified answer unit after its gates and session save. Provider token fragments and partial warnings remain withheld. This is a delivery contract; it does not promise that repair text appears during selection. C is not an automatic fallback. These choices are prospective UI methods, not a claim that the public app is deployed.')

    n.add_page_break();heading(n,'Measured comparison and its limits')
    paragraph(n,'The latest hosted experiment used Gemma 26B gemma-4-26b-a4b-it with minimal thinking, the same manual/index and three cases: GEN003 mirrors, EXP003 signal cancellation and GEN026 clock settings. Nine empty-cache turns and nine immediate repeats returned 8/9 SUPPORTED in each pass. All A/B cases passed; C clock failed twice. There were 27 attempts, zero retries, 26 known HTTP200 results and one unknown HTTP status.')
    labels={'GEN003':'Mirrors','EXP003':'Signal','GEN026':'Clock'}
    rows=[]
    for r in sorted(REPORT['initial_rows'],key=lambda x:(x['architecture'],list(labels).index(x['record_id']))):
        t=r['timing'];c=r['client'];fmt=lambda x:'--' if x is None else f'{x:.2f}'
        rows.append([r['architecture'][0]+' '+labels[r['record_id']],fmt(t['retrieval_wall_seconds']),fmt(t['generator_provider_ttft_seconds']),fmt(c.get('first_verified_output_seconds')),fmt(c.get('complete_response_seconds')),fmt(t['pacing_idle_only_seconds']),fmt(r.get('client_full_without_idle_pacing_arithmetic_seconds'))])
    table(n,['Method case','Retrieval','TTFT','TTFO','Full','Pacing','Full minus pacing'],rows,[1.15,.85,.65,.7,.7,.85,1.5])
    paragraph(n,'All values are seconds. Retrieval includes assembly and conditional reranking. TTFT above begins at provider transport start and represents the first non-thinking text fragment, not safe advice. TTFO and full completion are actual localhost HTTP/SSE client receipt times from request start, including setup, retrieval and pacing. Request-relative TTFT is retained in the full JSON report. Public network and browser rendering are unmeasured.')
    paragraph(n,'Full minus pacing is arithmetic, not an observed unpaced run. Idle waits are counted by interval union, with concurrent active provider work excluded. Successful adjusted full times were approximately 1.80-4.84 seconds. Conservative pacing was an experimental admission policy, not a verified account quota; most long observed waits were this pacing.')
    paragraph(n,'C clock had no delivered answer. Initial client error time was 9.63 seconds; its exact transport/gate cause and HTTP status were not retained. The repeat failed at 49.16 seconds because both verifier payloads used a Boolean array instead of the required scalar unsupported_additions field. Invalid schema is a failed verification operation, not proof that the wording was unsupported.')

    n.add_page_break();heading(n,'Latency choices that were tested')
    table(n,['Design choice','Evidence from this run','Decision or interpretation'],[
        ['Sequential or concurrent dense and BM25','Three warm paired controls returned identical allowed envelopes. Concurrent first-stage time was 0.424-3.341 ms slower.','Concurrency remains implemented; no demonstrated latency benefit. BM25 took 0.30-0.70 ms, so embedding and reranking are larger opportunities.'],
        ['Candidate cache or bypass','All three A repeats hit; actual full times were 2.03, 1.24 and 1.18 seconds. B/C paraphrases missed.','Keep exact reviewed candidate reuse. Similarity is lookup only; scope/version/state bindings and fresh gates remain mandatory. Do not claim general semantic answer caching.'],
        ['Buffered or verified streaming','Mirrors and signal each required a complete safe unit. Only clock had two independent blocks, and C clock failed.','Select A streaming by owner decision. This run does not prove earlier TTFO from independent block streaming.'],
        ['Cold or warm process','Fresh worker initialization was 63.48 seconds, including 6.84 seconds of controls. Readiness excluding controls was 56.64 seconds.','Investigate packaging, warm reuse, minimum instances and startup CPU boost. No hosted configuration change has been made.']
    ],[1.55,2.65,2.9])
    paragraph(n,'Dense model/index loading consumed 38.25 seconds of startup, registry validation 2.88 seconds, reranker load 0.28 seconds and query warmups about 14.43 seconds. The separate pretrial cooldown was 60.00 seconds. GitHub installation and job setup are additional engineering overhead, not rider request latency. These measurements must not be relabeled as Cloud Run cold-request times.')
    paragraph(n,'The deployed Qdrant service separately restores stored immutable vectors on startup. It does not parse or re-embed the manual per request. CPU query-model loading in the benchmark worker and Qdrant restoration are distinct cold-start paths; each needs its own measurement.')
    paragraph(n,'This three-case engineering comparison supports the current bounded owner decision. It does not establish population p95, held-out safety performance, a universal speed ranking or field repair success. Reversing the choice requires comparable reviewed cases, clear failure accounting and unchanged evidence gates.')

    n.add_page_break();heading(n,'Broader design choices considered')
    paragraph(n,'The design work also considered the alternatives below. The earlier decision register in 04_To_Build_Evaluation_Decisions.docx and HOSTED_INDEX_DECISION.md contains the detailed history. Options marked considered were not all implemented or tested; the distinction prevents an alternatives list from implying evidence that does not exist.')
    table(n,['Area','Current design and alternatives','Evidence status'],[
        ['Source and parsing','Immutable owner manual; Docling with flagged page 103 fallback. Other layout parsers or Document AI considered.','Hosted ingestion completed; no approved comparative parser winner.'],
        ['Chunking','1,024-token parent-child chunks; flat chunks, late chunking and alternate procedure boundaries considered.','Selected candidate has 532 children and 277 parents; no held-out chunking superiority claim.'],
        ['Embedding','BGE M3 selected; MiniLM compared.','Draft hybrid cited-page coverage was BGE 23/27 versus MiniLM 21/27; in-sample diagnostic only.'],
        ['Retrieval','Dense plus BM25 with RRF; dense-only and sparse-only alternatives.','Hybrid implemented; exact relevance scores remain candidate signals, not support proof.'],
        ['Graph','Lightweight verified provenance links; full GraphRAG and wider hops considered.','Only verified dependencies can authorize a step; broad graph benefit remains unproven.'],
        ['Reranking','Conditional cross-encoder; off, always-diagnostic and late interaction alternatives.','Conditional policy owner approved. Earlier BGE page coverage 23/27 without versus 22/27 with reranking; calibration remains scope-bound.'],
        ['Controller','Deterministic typed state and TRUE/FALSE/UNKNOWN eligibility; Decisions API or Jev considered.','Models may interpret or rank eligible options, never replace controller authority.'],
        ['Grounding','Exact reviewed wording in A or composition plus verification in B; independent block verification in C.','A/B selected after measured experiment; post facto A audit remains a UI prerequisite.'],
        ['Photo and voice','Visible observations only; visual-page retrieval or ColPali considered. Optional STT, including Sarvam, considered.','General vision/voice performance not established by this text benchmark.'],
        ['Hosting and traces','Private hosted build, Cloud Run Qdrant and Storage; Streamlit UI. Managed vector services, other UI hosts and Langfuse considered.','Private retrieval deployment exists; public UI, cost and end-to-end delivery checks are separate work.']
    ],[1.1,3.45,2.55])

    n.add_page_break();heading(n,'Next work before UI integration')
    paragraph(n,'The owner requested a new chat to complete two concrete tasks. The durable handoff is TRD_UI_FOLLOWUP_HANDOFF.md. Do not rerun composition merely to test whether the old output passes verification, and do not enable C in the UI.')
    heading(n,'Post facto semantic audit of A')
    paragraph(n,'Recover the exact A wording, claim order, citations and scope from hosted run 38122505406 and its approved registry bindings. Run the removed verification step against the original cited source and required reviewed meaning for mirrors, signal and clock. Retain initial and repeat provenance; identical repeat wording may share one verifier result with an explicit exact hash match. The result is pending, not assumed to pass.')
    paragraph(n,'Record each claim as validly accepted, validly rejected or unassessed because of malformed schema/provider failure. Keep conditions, warning completeness, unsupported additions and citation support explicit. Use a fresh bounded hosted audit protocol, budget check and pacing ledger rather than reusing the exhausted v4 benchmark authorization. Keep Gemma 26B unless the owner changes it; never relax the scalar Boolean schema to hide a failure.')
    heading(n,'Minimum instance cost and startup CPU boost')
    paragraph(n,'Investigate current official Cloud Run prices for us-central1 and the actual service CPU, memory, billing mode and minimum-instance setting. Show zero versus one minimum instance under stated idle/traffic assumptions, including monthly and daily cost, free-tier/account dependence, INR assumptions and impact on the INR800 lifetime budget and INR640 billing stop. Distinguish the Qdrant service from any future query-model/API service.')
    paragraph(n,'Read the live startup CPU boost setting and startup path, check official pricing and propose a repeatable cold/warm experiment. Packaging pinned model assets, loading once and reusing a process are candidate improvements; verify the relevant service before attributing the measured 38.25-second worker load to Cloud Run. Research authorizes no deployment, minimum-instance, CPU, billing or public source-display change.')
    heading(n,'Evidence record')
    paragraph(n,'Hosted run: https://github.com/Maneshwar-Kumar/motomedic-ai/actions/runs/38122505406\nReviewed runtime commit: e910c86d2fb8a0a2f955bbea4fa4cbedbe64d82d\nRaw results: reports/architecture_run_38122505406/architecture_comparison.json\nTiming summary: reports/architecture_owner_measurement_summary_v4.json\nIndependent audit: reports/architecture_result_independent_audit_v4.json\nRaw SHA256: bfd7603c1dccb84d99613c8de647de9475da4f9675d75ccdfb583c915583ac83')
    paragraph(n,'Earlier v3 run 38080802638 and the original comparison reports remain preserved. The v4 independent audit confirms all 18 cells, provider accounting and client timing arithmetic; it does not certify broad rider release. Current minimum-instance and CPU-boost reference pages are https://docs.cloud.google.com/run/docs/configuring/min-instances and https://docs.cloud.google.com/run/docs/configuring/services/cpu. Pricing will be checked afresh in the follow-up.')
    n.save(ROOT/NOTE)

if __name__=='__main__':main()
